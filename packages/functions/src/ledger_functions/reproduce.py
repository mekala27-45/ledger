"""Has a rederived output reproduced its committed copy?

Documents must match byte for byte: they are rendered from rounded values, so any change
in one is a change in what a reader is told. Data files are compared by content. DuckDB
and Polars sum floats in parallel, in no fixed order, so a rerun can move a float in its
last bits; a number that agrees to a relative one part in a billion has reproduced, since
nothing is published to anything like that precision. Parquet files can also differ in
their bytes while holding the same rows, which a byte comparison would call a change.
"""

from __future__ import annotations

import io
import json
import math
from pathlib import Path
from typing import Any

import numpy as np
import polars as pl

TOLERANCE = 1e-9
FRAMES = {".parquet", ".arrow", ".csv"}


def values_close(a: Any, b: Any, tol: float = TOLERANCE) -> bool:
    """Structural equality, with any pair of numbers where one is a float compared at ``tol``."""
    if isinstance(a, bool) or isinstance(b, bool):
        return a is b
    if isinstance(a, int | float) and isinstance(b, int | float):
        if isinstance(a, int) and isinstance(b, int):
            return a == b
        if math.isnan(a) or math.isnan(b):
            return math.isnan(a) and math.isnan(b)
        return math.isclose(a, b, rel_tol=tol, abs_tol=tol)
    if isinstance(a, dict) and isinstance(b, dict):
        return a.keys() == b.keys() and all(values_close(a[k], b[k], tol) for k in a)
    if isinstance(a, list) and isinstance(b, list):
        return len(a) == len(b) and all(values_close(x, y, tol) for x, y in zip(a, b, strict=True))
    return bool(a == b)


def frames_close(a: pl.DataFrame, b: pl.DataFrame, tol: float = TOLERANCE) -> bool:
    """Same columns, types and rows in the same order; float columns compared at ``tol``."""
    if a.schema != b.schema or a.height != b.height:
        return False
    for name in a.columns:
        x, y = a[name], b[name]
        if x.dtype.is_float():
            if not x.is_null().equals(y.is_null()):
                return False
            xs, ys = x.drop_nulls().to_numpy(), y.drop_nulls().to_numpy()
            if not np.allclose(xs, ys, rtol=tol, atol=tol, equal_nan=True):
                return False
        elif not x.equals(y):
            return False
    return True


def reproduced(path: Path, committed: bytes, tol: float = TOLERANCE) -> bool:
    """Whether the file at ``path`` holds what ``committed`` held, by the rules above."""
    current = path.read_bytes()
    if current == committed:
        return True
    suffix = path.suffix.lower()
    if suffix in FRAMES:
        return frames_close(_frame(suffix, committed), _frame(suffix, current), tol)
    if suffix == ".json":
        return values_close(json.loads(committed), json.loads(current), tol)
    if suffix == ".xlsx":
        return values_close(_workbook(committed), _workbook(current), tol)
    return False


def _frame(suffix: str, data: bytes) -> pl.DataFrame:
    buf = io.BytesIO(data)
    if suffix == ".parquet":
        return pl.read_parquet(buf)
    if suffix == ".arrow":
        return pl.read_ipc(buf)
    return pl.read_csv(buf)


def _workbook(data: bytes) -> dict[str, list[list[Any]]]:
    """Every sheet's cells, formulas as written. The file's own timestamps are not content."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}
